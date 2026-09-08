# -*- coding: utf-8 -*-
"""Compare DailyObservations CSV vs HSE TEAM xlsx for Mohamed Saeed."""
import csv
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

CSV_PATH = Path(r"C:\Users\YasserMohamed\Downloads\V.3-HSE Database - DailyObservations.csv")
XLSX_PATH = Path(r"C:\Users\YasserMohamed\Downloads\Safety Observation - HSE TEAM  (1).xlsx")
OUT = Path(r"d:\Apps\2026-07\clinic-repo\scripts\_obs-compare-mohamed-saeed.json")


def norm(s):
    s = "" if s is None else str(s)
    s = s.replace("\u0640", "")
    s = re.sub(r"\s+", " ", s).strip().lower()
    trans = str.maketrans({
        "أ": "ا", "إ": "ا", "آ": "ا", "ى": "ي", "ة": "ه", "ؤ": "و", "ئ": "ي",
    })
    return s.translate(trans)


def name_match(s):
    n = norm(s)
    if not n:
        return False
    compact = n.replace(" ", "")
    has_mohamed = ("محمد" in n) or ("mohamed" in n) or ("mohammed" in n)
    has_saeed = ("سعيد" in n) or ("saeed" in n) or ("said" in n)
    has_sharaf = ("شرف" in n) or ("sharaf" in n)
    has_hassan = ("حسن" in n) or ("hassan" in n) or ("hasan" in n)
    if has_mohamed and has_saeed and (has_sharaf or has_hassan):
        return True
    if "محمدسعيد" in compact and ("شرف" in compact or "حسنشرف" in compact):
        return True
    return False


def main():
    # ---- CSV ----
    observers = Counter()
    csv_obs = []
    with CSV_PATH.open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            on = row.get("observerName") or ""
            observers[on] += 1
            if name_match(on):
                csv_obs.append(row)

    print("=== CSV similar observerName ===")
    for name, cnt in observers.most_common():
        nn = norm(name)
        if any(k in nn for k in ("سعيد", "شرف", "saeed", "sharaf", "محمد سعيد")):
            print(f"  [{cnt:4d}] {name!r}")
    print(f"CSV observer matches: {len(csv_obs)}")

    # ---- XLSX ----
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)
    print("\nXLSX sheets:", wb.sheetnames)
    xlsx_matches = []
    xlsx_meta = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_raw = next(rows_iter)
        except StopIteration:
            continue
        header = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(header_raw)]
        print(f"\n--- Sheet {sheet_name!r} cols={len(header)} ---")
        print("HEADER:", " | ".join(header[:35]))

        name_idx = []
        for i, h in enumerate(header):
            hl = norm(h)
            if any(
                k in hl
                for k in (
                    "observer", "name", "reported", "employee", "inspector",
                    "ملاحظ", "راصد", "اسم", "submitted", "created by", "observername",
                    "hse", "officer", "team",
                )
            ):
                name_idx.append(i)
                print(" name-col", i, h)
        xlsx_meta[sheet_name] = {"header": header, "name_idx": name_idx}

        scanned = 0
        sheet_hits = 0
        for ridx, row in enumerate(rows_iter, start=2):
            scanned += 1
            vals = [("" if v is None else str(v)) for v in row]
            hit = False
            for i in name_idx:
                if i < len(vals) and name_match(vals[i]):
                    hit = True
                    break
            if not hit:
                # fallback: any cell (may be noisy)
                for v in vals:
                    if name_match(v):
                        hit = True
                        break
            if hit:
                sheet_hits += 1
                rec = {
                    (header[i] if i < len(header) else f"c{i}"): (vals[i] if i < len(vals) else "")
                    for i in range(len(header))
                }
                rec["_sheet"] = sheet_name
                rec["_row"] = ridx
                xlsx_matches.append(rec)
        print(f" rows scanned={scanned}, matches={sheet_hits}")

    wb.close()

    print("\n=== TOTAL ===")
    print("CSV:", len(csv_obs), "XLSX:", len(xlsx_matches))
    print("CSV names:", Counter((r.get("observerName") or "").strip() for r in csv_obs))

    by_sheet = defaultdict(list)
    for m in xlsx_matches:
        by_sheet[m["_sheet"]].append(m)
    for sh, items in by_sheet.items():
        meta = xlsx_meta[sh]
        print(f"\nSheet {sh}: {len(items)}")
        name_counter = Counter()
        for m in items:
            for i in meta["name_idx"]:
                key = meta["header"][i]
                val = (m.get(key) or "").strip()
                if val:
                    name_counter[val] += 1
        for n, c in name_counter.most_common(20):
            print(f"  [{c}] {n!r}")

    def slim_csv(r):
        keys = [
            "id", "isoCode", "siteName", "locationName", "observationType",
            "observerName", "date", "shift", "details", "correctiveAction",
            "status", "riskLevel", "createdAt", "updatedAt",
        ]
        return {k: r.get(k, "") for k in keys}

    def slim_xlsx(r):
        d = {k: v for k, v in r.items() if not str(k).startswith("_") and str(v).strip()}
        d["_sheet"] = r["_sheet"]
        d["_row"] = r["_row"]
        return d

    # date/status stats from CSV
    csv_dates = sorted({str(r.get("date") or "")[:10] for r in csv_obs if r.get("date")})
    csv_status = Counter((r.get("status") or "").strip() for r in csv_obs)
    csv_sites = Counter((r.get("siteName") or "").strip() for r in csv_obs)
    csv_types = Counter((r.get("observationType") or "").strip() for r in csv_obs)

    payload = {
        "csv_count": len(csv_obs),
        "xlsx_count": len(xlsx_matches),
        "csv_observer_names": dict(Counter((r.get("observerName") or "").strip() for r in csv_obs)),
        "csv_status": dict(csv_status),
        "csv_sites": dict(csv_sites),
        "csv_types": dict(csv_types),
        "csv_date_min": csv_dates[0] if csv_dates else None,
        "csv_date_max": csv_dates[-1] if csv_dates else None,
        "csv_date_count": len(csv_dates),
        "csv_all": [slim_csv(r) for r in csv_obs],
        "xlsx_all": [slim_xlsx(r) for r in xlsx_matches],
    }
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print("\nWrote", OUT)


if __name__ == "__main__":
    main()
