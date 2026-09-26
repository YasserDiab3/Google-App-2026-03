import os
import sys
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

PROJECT_ROOT = r"D:\Apps\1.9.2026\ICAPP.V.7.09-2026"
SCREENSHOTS_DIR = os.path.join(PROJECT_ROOT, "screenshots")
LOGO_PATH = os.path.join(PROJECT_ROOT, "Frontend", "icons", "icapp-logo.png")

# Palette
NAVY_DARK = RGBColor(10, 37, 64)       # #0A2540
NAVY_MID = RGBColor(24, 53, 88)        # #183558
ROYAL_BLUE = RGBColor(30, 136, 229)    # #1E88E5
LIGHT_BLUE = RGBColor(235, 244, 255)   # #EBF4FF
EMERALD = RGBColor(16, 149, 116)       # #109574
EMERALD_BG = RGBColor(236, 253, 245)   # #ECFDF5
AMBER = RGBColor(217, 119, 6)          # #D97706
AMBER_BG = RGBColor(254, 243, 199)     # #FEF3C7
CRIMSON = RGBColor(220, 38, 38)        # #DC2626
CRIMSON_BG = RGBColor(254, 242, 242)   # #FEF2F2
PURPLE = RGBColor(124, 58, 237)        # #7C3AED
PURPLE_BG = RGBColor(245, 243, 255)    # #F5F3FF
SLATE_DARK = RGBColor(15, 23, 42)      # #0F172A
SLATE_MUTED = RGBColor(100, 116, 139)  # #64748B
CARD_BG = RGBColor(255, 255, 255)      # White
PAGE_BG = RGBColor(248, 250, 252)      # #F8FAFC
BORDER_COLOR = RGBColor(226, 232, 240) # #E2E8F0

def add_header_footer(slide, category_text, title_text, subtitle_text, slide_num, total_slides, is_ar=True):
    # Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = PAGE_BG
    bg.line.fill.background()

    # Top Bar
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.15))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = NAVY_DARK
    top_bar.line.fill.background()

    # Accent Line
    acc = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.15), Inches(13.333), Inches(0.04))
    acc.fill.solid()
    acc.fill.fore_color.rgb = ROYAL_BLUE
    acc.line.fill.background()

    # Category Pill
    pill_x = Inches(9.8) if is_ar else Inches(0.6)
    cat_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pill_x, Inches(0.2), Inches(2.9), Inches(0.32))
    cat_box.fill.solid()
    cat_box.fill.fore_color.rgb = NAVY_MID
    cat_box.line.color.rgb = ROYAL_BLUE
    cat_box.line.width = Pt(1)
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    p_c = tf_c.paragraphs[0]
    p_c.text = category_text
    p_c.alignment = PP_ALIGN.CENTER
    p_c.font.size = Pt(11)
    p_c.font.bold = True
    p_c.font.color.rgb = RGBColor(186, 230, 253)
    p_c.font.name = "Segoe UI"

    # Title & Subtitle Box
    t_x = Inches(3.2) if is_ar else Inches(3.7)
    title_box = slide.shapes.add_textbox(t_x, Inches(0.12), Inches(6.4), Inches(0.55))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.alignment = PP_ALIGN.RIGHT if is_ar else PP_ALIGN.LEFT
    p_t.font.size = Pt(19)
    p_t.font.bold = True
    p_t.font.color.rgb = RGBColor(255, 255, 255)
    p_t.font.name = "Segoe UI"

    sub_box = slide.shapes.add_textbox(t_x, Inches(0.62), Inches(6.4), Inches(0.45))
    tf_s = sub_box.text_frame
    tf_s.word_wrap = True
    p_s = tf_s.paragraphs[0]
    p_s.text = subtitle_text
    p_s.alignment = PP_ALIGN.RIGHT if is_ar else PP_ALIGN.LEFT
    p_s.font.size = Pt(11)
    p_s.font.color.rgb = RGBColor(203, 213, 225)
    p_s.font.name = "Segoe UI"

    # Logo
    logo_x = Inches(0.5) if is_ar else Inches(11.8)
    if os.path.exists(LOGO_PATH):
        try:
            slide.shapes.add_picture(LOGO_PATH, logo_x, Inches(0.16), height=Inches(0.82))
        except Exception:
            pass

    # Footer
    f_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.05), Inches(13.333), Inches(0.02))
    f_line.fill.solid()
    f_line.fill.fore_color.rgb = BORDER_COLOR
    f_line.line.fill.background()

    footer_box = slide.shapes.add_textbox(Inches(0.6), Inches(7.08), Inches(12.133), Inches(0.35))
    tf_f = footer_box.text_frame
    p_f = tf_f.paragraphs[0]
    if is_ar:
        p_f.text = f"SafetyHub | ICAPP  •  نظام إدارة السلامة والصحة المهنية والبيئة  •  شريحة {slide_num} من {total_slides}"
    else:
        p_f.text = f"SafetyHub | ICAPP  •  HSE Management System  •  Slide {slide_num} of {total_slides}"
    p_f.alignment = PP_ALIGN.CENTER
    p_f.font.size = Pt(9.5)
    p_f.font.color.rgb = SLATE_MUTED
    p_f.font.name = "Segoe UI"

def add_visual_content(slide, image_filename, card_title, bullets, impact_badge, quick_specs=None, is_ar=True, accent_color=ROYAL_BLUE):
    if is_ar:
        text_left = Inches(6.9)
        img_container_left = Inches(0.6)
    else:
        text_left = Inches(0.6)
        img_container_left = Inches(6.9)

    top_y = Inches(1.3)
    card_w = Inches(5.8)
    card_h = Inches(5.55)

    # 1. Text Card
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, text_left, top_y, card_w, card_h)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = BORDER_COLOR
    card.line.width = Pt(1)

    # Card Banner
    banner = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, text_left, top_y, card_w, Inches(0.5))
    banner.fill.solid()
    banner.fill.fore_color.rgb = accent_color
    banner.line.fill.background()

    tb = slide.shapes.add_textbox(text_left + Inches(0.15), top_y + Inches(0.06), card_w - Inches(0.3), Inches(0.4))
    tf_tb = tb.text_frame
    p_tb = tf_tb.paragraphs[0]
    p_tb.text = card_title
    p_tb.alignment = PP_ALIGN.RIGHT if is_ar else PP_ALIGN.LEFT
    p_tb.font.size = Pt(13.5)
    p_tb.font.bold = True
    p_tb.font.color.rgb = RGBColor(255, 255, 255)
    p_tb.font.name = "Segoe UI"

    # Bullets
    bb = slide.shapes.add_textbox(text_left + Inches(0.2), top_y + Inches(0.6), card_w - Inches(0.4), Inches(4.25))
    tf_bb = bb.text_frame
    tf_bb.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf_bb.paragraphs[0] if i == 0 else tf_bb.add_paragraph()
        p.text = f"• {b}"
        p.alignment = PP_ALIGN.RIGHT if is_ar else PP_ALIGN.LEFT
        p.font.size = Pt(10.8)
        p.font.name = "Segoe UI"
        p.font.color.rgb = SLATE_DARK
        p.space_after = Pt(6)

    # Impact Badge
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, text_left + Inches(0.2), top_y + card_h - Inches(0.55), card_w - Inches(0.4), Inches(0.42))
    badge.fill.solid()
    badge.fill.fore_color.rgb = LIGHT_BLUE
    badge.line.color.rgb = ROYAL_BLUE
    badge.line.width = Pt(1)
    tf_badge = badge.text_frame
    tf_badge.word_wrap = True
    p_bg = tf_badge.paragraphs[0]
    p_bg.text = impact_badge
    p_bg.alignment = PP_ALIGN.CENTER
    p_bg.font.size = Pt(10.5)
    p_bg.font.bold = True
    p_bg.font.color.rgb = NAVY_DARK
    p_bg.font.name = "Segoe UI"

    # 2. Image Container
    img_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, img_container_left, top_y, card_w, card_h)
    img_card.fill.solid()
    img_card.fill.fore_color.rgb = CARD_BG
    img_card.line.color.rgb = BORDER_COLOR
    img_card.line.width = Pt(1)

    # Image header strip
    img_header = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, img_container_left, top_y, card_w, Inches(0.38))
    img_header.fill.solid()
    img_header.fill.fore_color.rgb = NAVY_MID
    img_header.line.fill.background()

    img_title_box = slide.shapes.add_textbox(img_container_left + Inches(0.15), top_y + Inches(0.04), card_w - Inches(0.3), Inches(0.3))
    tf_it = img_title_box.text_frame
    p_it = tf_it.paragraphs[0]
    p_it.text = "🖥️ لقطة شاشة حية من النظام الفعلي" if is_ar else "🖥️ Live System Screenshot"
    p_it.alignment = PP_ALIGN.CENTER
    p_it.font.size = Pt(10)
    p_it.font.bold = True
    p_it.font.color.rgb = RGBColor(224, 242, 254)
    p_it.font.name = "Segoe UI"

    # Add Image
    img_path = os.path.join(SCREENSHOTS_DIR, image_filename)
    if os.path.exists(img_path):
        try:
            img_w = Inches(5.5)
            img_h = Inches(3.45)
            img_x = img_container_left + Inches(0.15)
            img_y = top_y + Inches(0.48)
            slide.shapes.add_picture(img_path, img_x, img_y, width=img_w, height=img_h)
        except Exception as e:
            print(f"Error adding picture {img_path}: {e}")

    # Bottom Quick Specs inside Image Container
    specs_y = top_y + Inches(4.02)
    specs_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, img_container_left + Inches(0.15), specs_y, card_w - Inches(0.3), Inches(1.38))
    specs_box.fill.solid()
    specs_box.fill.fore_color.rgb = RGBColor(241, 245, 249)
    specs_box.line.color.rgb = BORDER_COLOR
    specs_box.line.width = Pt(1)

    tf_sp = specs_box.text_frame
    tf_sp.word_wrap = True
    tf_sp.margin_left = tf_sp.margin_right = tf_sp.margin_top = tf_sp.margin_bottom = Inches(0.08)

    sp_title = tf_sp.paragraphs[0]
    sp_title.text = "📊 أرقام وحقائق حية من قاعدة بيانات النظام:" if is_ar else "📊 Real Database Metrics & Field Facts:"
    sp_title.alignment = PP_ALIGN.RIGHT if is_ar else PP_ALIGN.LEFT
    sp_title.font.size = Pt(10)
    sp_title.font.bold = True
    sp_title.font.color.rgb = NAVY_DARK
    sp_title.font.name = "Segoe UI"
    sp_title.space_after = Pt(2)

    specs_list = quick_specs if quick_specs else []
    for spec in specs_list:
        p_sp = tf_sp.add_paragraph()
        p_sp.text = f"✔ {spec}"
        p_sp.alignment = PP_ALIGN.RIGHT if is_ar else PP_ALIGN.LEFT
        p_sp.font.size = Pt(9.2)
        p_sp.font.color.rgb = SLATE_DARK
        p_sp.font.name = "Segoe UI"

def generate_arabic_presentation():
    print("Generating Arabic Presentation with REAL Database Numbers...")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    TOTAL_SLIDES = 19

    # SLIDE 1: Cover
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY_DARK
    bg1.line.fill.background()

    acc1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.6), Inches(12.133), Inches(6.3))
    acc1.fill.background()
    acc1.line.color.rgb = ROYAL_BLUE
    acc1.line.width = Pt(1.5)

    if os.path.exists(LOGO_PATH):
        try:
            s1.shapes.add_picture(LOGO_PATH, Inches(5.66), Inches(1.1), height=Inches(1.2))
        except Exception:
            pass

    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(2.45), Inches(11.333), Inches(1.1))
    p1 = t_box.text_frame.paragraphs[0]
    p1.text = "SafetyHub | ICAPP"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)
    p1.font.name = "Segoe UI"

    s_box = s1.shapes.add_textbox(Inches(1.0), Inches(3.55), Inches(11.333), Inches(0.6))
    ps1 = s_box.text_frame.paragraphs[0]
    ps1.text = "العرض التقديمي المصور لمنظومة السلامة والصحة المهنية والبيئة (بيانات حية)"
    ps1.alignment = PP_ALIGN.CENTER
    ps1.font.size = Pt(21)
    ps1.font.bold = True
    ps1.font.color.rgb = RGBColor(125, 211, 252)
    ps1.font.name = "Segoe UI"

    s_box2 = s1.shapes.add_textbox(Inches(1.0), Inches(4.25), Inches(11.333), Inches(0.5))
    ps2 = s_box2.text_frame.paragraphs[0]
    ps2.text = "الشركة العالمية للإنتاج والتصنيع الزراعي (ICAPP) • توثيق ميداني مدعوم بأرقام وسجلات قاعدة البيانات"
    ps2.alignment = PP_ALIGN.CENTER
    ps2.font.size = Pt(15)
    ps2.font.color.rgb = RGBColor(226, 232, 240)
    ps2.font.name = "Segoe UI"

    badge_titles = [
        "📊 1,654 تصريح عمل مسجل",
        "📋 3,113 ملاحظة ميدانية موثقة",
        "🏥 1,029 زيارة عيادة مسجلة",
        "🔒 41,744 سجل تدقيق أمني"
    ]
    bw = Inches(2.7)
    b_gap = Inches(0.25)
    start_bx = Inches(1.0)
    for i, b_text in enumerate(badge_titles):
        bx = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, start_bx + i * (bw + b_gap), Inches(5.1), bw, Inches(0.55))
        bx.fill.solid()
        bx.fill.fore_color.rgb = NAVY_MID
        bx.line.color.rgb = ROYAL_BLUE
        bx.line.width = Pt(1)
        p_bx = bx.text_frame.paragraphs[0]
        p_bx.text = b_text
        p_bx.alignment = PP_ALIGN.CENTER
        p_bx.font.size = Pt(11)
        p_bx.font.bold = True
        p_bx.font.color.rgb = RGBColor(255, 255, 255)
        p_bx.font.name = "Segoe UI"

    f_box = s1.shapes.add_textbox(Inches(1.0), Inches(6.0), Inches(11.333), Inches(0.5))
    p_f1 = f_box.text_frame.paragraphs[0]
    p_f1.text = "إدارة السلامة والصحة المهنية والبيئة — شركة ICAPP • تم التحديث ببيانات الإنتاج الحية © 2026"
    p_f1.alignment = PP_ALIGN.CENTER
    p_f1.font.size = Pt(11)
    p_f1.font.color.rgb = RGBColor(148, 163, 184)
    p_f1.font.name = "Segoe UI"

    # SLIDE 2: Login AR
    s2 = prs.slides.add_slide(blank_layout)
    add_header_footer(s2, "الأمان والمصادقة", "بوابة تسجيل الدخول والأمان والمصادقة الثنائية", "واجهة دخول آمنة متعددة اللغات مع تشفير تام وإدارة جلسات موثوقة", 2, TOTAL_SLIDES, is_ar=True)
    add_visual_content(s2, "01_login_realistic.png", "مميزات بوابة الدخول المعتمدة", [
        "تسجيل دخول آمن بالبريد المؤسسي مع تشفير SHA-256 لكلمات المرور.",
        "دعم المصادقة الثنائية (2FA/MFA) عبر تطبيقات Authenticator المعتمدة.",
        "إمكانية الدخول السريع برمز PIN للمشرفين الميدانيين لتوفير الوقت.",
        "تبديل فوري بين اللغتين العربية والإنجليزية مع تعديل تلقائي للاتجاه (RTL/LTR).",
        "حماية مشددة ضد محاولات الاختراق مع تسجيل كامل لـ 41,744 سجل تدقيق أمني."
    ], "الأمان: حماية سيبرانية مشددة تمنع أي وصول غير مصرح به", [
        "تسجيل وتدقيق 41,744 عملية أمنية في SecurityAuditLog.",
        "توثيق 6,470 جلسة ونشاط مستخدم في UserActivityLog.",
        "إدارة انتهاء الجلسات التلقائي وتشفير التوكنات السحابية."
    ], is_ar=True, accent_color=ROYAL_BLUE)

    # SLIDE 3: Dashboard AR
    s3 = prs.slides.add_slide(blank_layout)
    add_header_footer(s3, "لوحة القيادة", "لوحة التحكم التنفيذية والمؤشرات اللحظية (KPIs)", "نظرة شمولية فورية على مدار 24 ساعة لكافة مؤشرات وأنشطة السلامة", 3, TOTAL_SLIDES, is_ar=True)
    add_visual_content(s3, "03_dashboard_realistic.png", "وظائف لوحة التحكم الرئيسية", [
        "بطاقات إحصائية لحظية: ساعات العمل الآمنة، التصاريح النشطة، ونسب الإغلاق.",
        "متابعة 57 مؤشر أداء رئيسي للسلامة والصحة المهنية (Safety KPIs).",
        "رسوم بيانية تفاعلية توضح التوزيع الجغرافي للمخاطر وتصنيفها الموقعي.",
        "شريط التنبيهات الذكي لرصد الحالات الحرجة ومواعيد الفحوصات المنتهية.",
        "أزرار وصول سريع لأكثر من 132 موقعاً وقسماً مسجلاً في النظام."
    ], "القيمة: رؤية بيانات فورية تدعم اتخاذ القرارات الوقائية الاستباقية", [
        "متابعة 57 مؤشر أداء سلامة معتمد في SafetyPerformanceKPIs.",
        "تغطية جغرافية تفصيلية لـ 132 موقع عمل في Form_Places.",
        "تحديث حي للبيانات دون الحاجة لتحديث الصفحة يدوياً."
    ], is_ar=True, accent_color=EMERALD)

    # SLIDE 4: PTW AR
    s4 = prs.slides.add_slide(blank_layout)
    add_header_footer(s4, "العمليات الحرجة", "منظومة تصاريح العمل الإلكترونية (e-PTW)", "إصدار واعتماد رقمي صارم للسيطرة الكاملة على الأعمال عالية الخطورة", 4, TOTAL_SLIDES, is_ar=True)
    add_visual_content(s4, "04_ptw_realistic.png", "دورة عمل تصاريح العمل الإلكترونية", [
        "إدارة وأرشفة 1,654 تصريح عمل إلكتروني فعلي و 1,941 حركة بسجل PTWRegistry.",
        "شبكة اعتماد مكونة من 63 مسؤول توقيع معتمد و 29 مسؤول للمقاولين.",
        "تغطية كاملة للأنشطة الخطرة: الساخن، الأماكن المغلقة، الارتفاعات، والعزل.",
        "توقيعات إلكترونية مشفرة مع توليد رمز QR فريد للتفتيش الميداني السريع.",
        "إغلاق نظامي مؤتمت مع التحقق من نظافة الموقع وإلغاء التصاريح عند الطوارئ."
    ], "النتيجة: 1,654 تصريح عمل رقمي تم تنفيذه دون أي حادث جسيم", [
        "1,654 تصريح عمل إلكتروني موثق في قاعدة بيانات PTW.",
        "1,941 حركة وسجل تتبع في جدول PTWRegistry.",
        "63 مسؤول توقيع داخلي + 29 مسؤول للمقاولين معتمدين."
    ], is_ar=True, accent_color=CRIMSON)

    # SLIDE 5: Incidents & RCA AR
    s5 = prs.slides.add_slide(blank_layout)
    add_header_footer(s5, "الحوادث والتحقيقات", "إدارة الحوادث وتحليل الأسباب الجذرية (RCA)", "توثيق شامل ومنهجيات تحقيق عالمية تمنع تكرار الحوادث والإصابات", 5, TOTAL_SLIDES, is_ar=True)
    add_visual_content(s5, "05_incidents_realistic.png", "إجراءات التعامل مع الحوادث والتحقيق", [
        "تسجيل فوري للحوادث مع إرفاق الصور والأدلة وتصنيف شدة الإصابة.",
        "أدوات تحقيق مدمجة: منهجية 'لماذا الخمسة' (5-Whys) ومخطط عظمة السمكة.",
        "تحديد العوامل الجذرية: الإنسان، الآلة، البيئة، الطريقة، والمواد.",
        "خطة إجراءات تصحيحية ووقائية (CAPA) مرتبطة بمسؤول وتاريخ إغلاق ملزم.",
        "ربط سجل الحوادث المباشر مع سجل تتبع الإجراءات (262 إجراء مسجل)."
    ], "الأثر: معالجة الأسباب العميقة لضمان عدم تكرار الحادث نهائياً", [
        "ربط مباشر بسجل المتابعة العام المكون من 262 إجراء (ATR).",
        "تتبع الحوادث والإصابات الطبية المحولة لعيادة المصنع.",
        "تقارير مطابقة لمعايير وزارة العمل والتأمينات الاجتماعية."
    ], is_ar=True, accent_color=CRIMSON)

    # SLIDE 6: Near-Miss AR
    s6 = prs.slides.add_slide(blank_layout)
    add_header_footer(s6, "الوقاية الاستباقية", "إدارة الحوادث الوشيكة (Near-Miss Reporting)", "رصد الإشارات المبكرة ومعالجة الظروف الخطرة قبل تحولها لحوادث فعلية", 6, TOTAL_SLIDES, is_ar=True)
    add_visual_content(s6, "06_nearmiss_realistic.png", "مميزات نظام الحوادث الوشيكة", [
        "شاشة إبلاغ مبسطة تمكن العمال من الإبلاغ السريع عبر الهاتف.",
        "تصنيف الخطر الكامن وتحديد أولويات المعالجة الفورية بالموقع.",
        "إمكانية الإبلاغ السري لتشجيع ثقافة الشفافية وكسر حاجز الخوف.",
        "تحليل تكرار المواقف الخطرة بالأقسام لتكثيف برامج الحماية الوقائية.",
        "متابعة إغلاق كافة البلاغات بنسبة 100% بالتعاون مع مسؤولي المواقع."
    ], "المبدأ: كل حادث وشيك يتم إغلاقه هو حادث جسيم تم منعه بنجاح", [
        "ربط بلاغات Near-Miss مباشرة مع مهام مسؤولي السلامة.",
        "إمكانية إرفاق صورة مباشرة من كاميرا الهاتف المحمول.",
        "تحليل استباقي يمنع تصاعد الأخطار البسيطة لأضرار بشرية."
    ], is_ar=True, accent_color=AMBER)

    # SLIDE 7: Clinic AR
    s7 = prs.slides.add_slide(blank_layout)
    add_header_footer(s7, "الصحة المهنية", "العيادة الطبية وبرامج الرعاية الصحية", "سجلات طبية رقمية مشفرة ومتابعة دقيقة لصحة ولياقة الكوادر البشرية", 7, TOTAL_SLIDES, is_ar=True)
    add_visual_content(s7, "07_clinic_realistic.png", "خدمات موديول العيادة الطبية", [
        "تسجيل 1,029 زيارة طبية للعمال والمقاولين في ClinicContractorVisits.",
        "توثيق 573 حركة صرف أدوية ومستلزمات في MedicationDispenseLog.",
        "سجل حضور الأطباء والتمريض يضم 187 سجلاً في ClinicStaffAttendance.",
        "إدارة الإجازات المرضية وساعات العمل الضائعة ومتابعة الحالات المزمنة.",
        "الفحوصات الدورية للعاملين في خطوط تصنيع الأغذية والبيئات الكيميائية."
    ], "الرعاية: 1,029 زيارة و 573 صرف دواء موثقة بالكامل في النظام", [
        "1,029 زيارة فحص وعلاج موثقة في ClinicContractorVisits.",
        "573 عملية صرف علاج موثقة في MedicationDispenseLog.",
        "187 سجل حضور لفريق التمريض والعيادة الطبية."
    ], is_ar=True, accent_color=EMERALD)

    # SLIDE 8: Fire Equipment AR
    s8 = prs.slides.add_slide(blank_layout)
    add_header_footer(s8, "مكافحة الحرائق", "فحص معدات الحريق ومكافحة الطوارئ", "جاهزية كاملة بنسبة 100% لكافة أجهزة وشبكات الإطفاء بالمنشأة", 8, TOTAL_SLIDES, is_ar=True)
    add_visual_content(s8, "08_fire_equipment_realistic.png", "منظومة إدارة معدات الإطفاء", [
        "حصر وتسجيل 169 أصل ومعدة إطفاء مسجلة في جدول FireEquipmentAssets.",
        "تنفيذ 26 دورة تفتيش دورية شاملة موثقة في FireEquipmentInspections.",
        "فحص ميداني فائق السرعة عبر مسح رمز QR Code الخاص بكل جهاز إطفاء.",
        "التحقق من ضغط العداد، سلامة الصمام، الخرطوم، وصلاحية بودرة الإطفاء.",
        "تنبيهات استباقية لمواعيد إعادة التعبئة والاختبارات الهيدروستاتيكية الدورية."
    ], "الجاهزية: 169 معدة إطفاء مفحوصة وموزعة جغرافياً في المصنع", [
        "169 أصل ومعدة إطفاء مسجلة في FireEquipmentAssets.",
        "26 دورة فحص وتفتيش شاملة في FireEquipmentInspections.",
        "جاهزية تامة وتوافق بنسبة 100% مع معايير الدفاع المدني."
    ], is_ar=True, accent_color=CRIMSON)

    # SLIDE 9: Contractors AR
    s9 = prs.slides.add_slide(blank_layout)
    add_header_footer(s9, "المقاولين والشركاء", "إدارة المقاولين، الزوار، وأمن البوابات", "حوكمة العمليات الخارجية والسيطرة على حركة الدخول والخروج بأعلى أمان", 9, TOTAL_SLIDES, is_ar=True)
    add_visual_content(s9, "09_contractors_realistic.png", "إجراءات حوكمة المقاولين والزوار", [
        "تدريب وتأهيل 1,736 عامل مقاول موثقين في جدول ContractorTrainings.",
        "معالجة 194 طلب اعتماد مقاولين و 56 طلب اعتماد تقييم أداء.",
        "قائمة سوداء رقمية تضم 13 مقاولاً ومخالفاً غير ملتزم في Blacklist_Register.",
        "تسجيل 29 حركة دخول زوار وشاحنات عبر البوابات في GateVisitors.",
        "إحصائيات لحظية للأفراد المتواجدين داخل المنشأة لضمان الإخلاء الآمن."
    ], "الحوكمة: 1,736 عامل مقاول مدرب و 194 طلب اعتماد رسمي", [
        "1,736 عامل وفني مقاول تم تدريبهم واختبارهم بنجاح.",
        "194 طلب اعتماد مقاول في ContractorApprovalRequests.",
        "13 مخالف في القائمة السوداء لحماية بيئة العمل من التجاوزات."
    ], is_ar=True, accent_color=ROYAL_BLUE)

    # SLIDE 10: PPE AR
    s10 = prs.slides.add_slide(blank_layout)
    add_header_footer(s10, "مهمات الوقاية", "إدارة مهمات الوقاية الشخصية (PPE) والمخزون", "خط الدفاع الأخير لحماية العاملين مع ترشيد وحوكمة الصرف والاستهلاك", 10, TOTAL_SLIDES, is_ar=True)
    add_visual_content(s10, "10_ppe_realistic.png", "منظومة مهمات الوقاية الشخصية", [
        "إدارة 162 سجلاً لمهمات الوقاية و 105 حركة صرف في PPE_Transactions.",
        "مصفوفة أهلية واستحقاق تغطي 12 قسماً ووظيفة في PPEMatrix.",
        "تحديد العمر الافتراضي لكل مهمة ومنع الصرف المتكرر غير المبرر.",
        "متابعة لحظية لرصيد مستودع مهمات الوقاية وتنبيهات إعادة الطلب.",
        "توقيع إلكتروني للموظف على استلام المهمات وتعهده بارتدائها الإلزامي."
    ], "الترشيد: 162 صنف مهمات و 105 حركة صرف موثقة بالكامل", [
        "162 سجلاً لمهمات الوقاية في جدول PPE.",
        "105 عملية صرف فردية ومجمعة في PPE_Transactions.",
        "12 مصفوفة استحقاق مهنية في PPEMatrix و PPE_Eligibility_Rules."
    ], is_ar=True, accent_color=AMBER)

    # SLIDE 11: Chemical Safety AR
    s11 = prs.slides.add_slide(blank_layout)
    add_header_footer(s11, "المواد الخطرة", "إدارة السلامة الكيميائية وصحائف الأمان (MSDS)", "السيطرة الشاملة على الكيماويات والتخزين الآمن واحتواء الانسكابات", 11, TOTAL_SLIDES, is_ar=True)
    add_visual_content(s11, "11_chemical_safety_realistic.png", "إجراءات السلامة الكيميائية", [
        "سجل مركزي للمواد الكيميائية مسجل في Chemical_Register مع أرقام CAS.",
        "أرشفة رقمية لصحائف بيانات الأمان (MSDS) باللغتين لسرعة الرجوع إليها.",
        "مصفوفة التوافق الكيميائي لمنع تخزين المواد المتفاعلة معاً في مكان واحد.",
        "إرشادات واضحة للإسعافات الأولية ومهمات الوقاية الخاصة بكل مادة.",
        "خطة احتواء الانسكابات وتأمين أحواض الاحتواء وأطقم المكافحة (Spill Kits)."
    ], "الوقاية: حماية العمال والبيئة من مخاطر التسمم والحرائق الكيميائية", [
        "حصر رسمي معتمد لكافة المواد الكيميائية في Chemical_Register.",
        "الترميز التحذيري العالمي GHS على جميع العبوات والمخازن.",
        "جاهزية كاملة لأطقم مكافحة الانسكاب ومحطات غسيل العيون."
    ], is_ar=True, accent_color=ROYAL_BLUE)

    # SLIDE 12: Periodic Inspections AR
    s12 = prs.slides.add_slide(blank_layout)
    add_header_footer(s12, "التفتيش الدوري", "الفحوصات الدورية للمعدات والمنشآت", "صيانة وقائية دورية تضمن سلامة الآلات، الرافعات، والأنظمة الكهربائية", 12, TOTAL_SLIDES, is_ar=True)
    add_visual_content(s12, "12_periodic_inspections_realistic.png", "نطاق الفحوصات الدورية المجدولة", [
        "تنفيذ وتوثيق الفحوصات الدورية المجدولة في PeriodicEquipmentInspections.",
        "قوائم فحص ذكية (Checklists) متوافقة مع الأكواد الهندسية المعتمدة.",
        "إلزامية التقاط صور للأعطال والعيوب المرصودة أثناء التفتيش الميداني.",
        "تحويل الملاحظات والعيوب تلقائياً إلى أوامر إصلاح فورية لفرق الصيانة.",
        "سجل تاريخي لكل آلة يوضح سجل أعطالها وتكلفة صيانتها السابقة."
    ], "الموثوقية: تفتيش هندسي مستمر يحافظ على سلامة الأصول والمعدات", [
        "توثيق عمليات الفحص الدوري في PeriodicInspections.",
        "فحص شامل لرافعات الشوكة، الغلايات، واللوحات الكهربائية.",
        "إصدار أوامر صيانة وقائية آلية لمنع التوقفات غير المخططة."
    ], is_ar=True, accent_color=EMERALD)

    # SLIDE 13: Violations AR
    s13 = prs.slides.add_slide(blank_layout)
    add_header_footer(s13, "الانضباط الميداني", "إدارة مخالفات السلامة ولائحة الجزاءات", "تطبيق عادل وصارم لقواعد السلامة لحماية بيئة العمل من السلوكيات الخاطئة", 13, TOTAL_SLIDES, is_ar=True)
    add_visual_content(s13, "13_violations_realistic.png", "حوكمة تسجيل وتطبيق المخالفات", [
        "تسجيل 154 مخالفة سلامة ميدانية فعلية موثقة في جدول Violations.",
        "معالجة 73 طلب اعتماد ومراجعة جزاءات في ViolationApprovalRequests.",
        "تصنيف المخالفات وفق 36 نموذجاً ولائحة معتمدة في جدول ViolationTypes.",
        "سلم جزاءات تصاعدي عادل: تنبيه شفهي -> لفت نظر -> إنذار -> جزاء مالي.",
        "ربط سجل المخالفات بنظام تقييم الأداء السنوي للموظفين والمقاولين."
    ], "الانضباط: 154 مخالفة مسجلة و 73 طلب مراجعة تم حسمها بعدالة", [
        "154 مخالفة ميدانية موثقة بالكامل في جدول Violations.",
        "73 طلب اعتماد إداري في ViolationApprovalRequests.",
        "36 نوع وتصنيف معتمد للمخالفات في ViolationTypes."
    ], is_ar=True, accent_color=AMBER)

    # SLIDE 14: Risk Assessment AR
    s14 = prs.slides.add_slide(blank_layout)
    add_header_footer(s14, "إدارة المخاطر", "تقييم المخاطر، مصفوفة 5x5، وتعليمات SOP/JHA", "تحديد استباقي لكافة الأخطار وتطبيق هرم التحكم لمنع وقوع الإصابات", 14, TOTAL_SLIDES, is_ar=True)
    add_visual_content(s14, "14_risk_assessment_realistic.png", "أدوات تقييم المخاطر المدمجة", [
        "مصفوفة تقييم المخاطر التفاعلية (5×5 Matrix) لحساب الاحتمالية والشدة.",
        "التحقق من 40 وثيقة ومستند امتثال قانوني في جدول LegalDocuments.",
        "تحليل سلامة الوظائف (JHA/JSA) بتفكيك المهام لخطوات وتحديد ضوابطها.",
        "مكتبة إجراءات التشغيل القياسية (SOPs) لسلامة تشغيل كافة الآلات.",
        "تطبيق هرم التحكم: الإزالة -> الاستبدال -> العزل الهندسي -> الإداري -> PPE."
    ], "الاستباق: السيطرة على مصادر الخطر والامتثال لـ 40 وثيقة قانونية", [
        "40 وثيقة امتثال قانوني مسجلة في جدول LegalDocuments.",
        "ربط مباشر بين تقييم المخاطر وتصاريح العمل الساخنة والمغلقة.",
        "إرشادات مصورة وواضحة يسهل فهمها من قبل العمال في الميدان."
    ], is_ar=True, accent_color=ROYAL_BLUE)

    # SLIDE 15: Training AR
    s15 = prs.slides.add_slide(blank_layout)
    add_header_footer(s15, "التدريب والتأهيل", "إدارة التدريب والسلامة المعرفية للكوادر", "بناء رأس المال البشري وترسيخ ثقافة الوعي بالسلامة لجميع المستويات", 15, TOTAL_SLIDES, is_ar=True)
    add_visual_content(s15, "15_training_realistic.png", "برامج التدريب والتأهيل المستمر", [
        "تنفيذ 439 دورة تدريبية داخلية موثقة في جدول Training.",
        "تسجيل 1,027 حضور موظف وتدريب 1,736 عامل مقاول في ContractorTrainings.",
        "مصفوفة تدريب تغطي 158 وظيفة وتخصص في EmployeeTrainingMatrix.",
        "إصدار شهادات التدريب الرقمية مزودة برمز تحقق QR Code معتمد.",
        "مكتبة محاضرات ما قبل العمل اليومية (Toolbox Talks - TBT) للمشرفين."
    ], "الكفاءة: 439 دورة تدريبية و 2,763 متدرب موثقين في السجلات", [
        "439 دورة تدريبية داخلية مسجلة في جدول Training.",
        "1,027 حضور موظف في TrainingAttendance + 1,736 مقاول.",
        "158 مصفوفة تدريب مهنية في EmployeeTrainingMatrix."
    ], is_ar=True, accent_color=EMERALD)

    # SLIDE 16: Forms Hub AR
    s16 = prs.slides.add_slide(blank_layout)
    add_header_footer(s16, "النماذج الميدانية", "الملاحظات اليومية وبوابة النماذج (Forms Hub)", "تمكين المشرفين والموظفين من الرصد الميداني السريع عبر الهاتف بدون إنترنت", 16, TOTAL_SLIDES, is_ar=True)
    add_visual_content(s16, "16_forms_hub_realistic.png", "بوابة النماذج الميدانية الميسرة", [
        "تسجيل وتوثيق 3,113 ملاحظة سلامة يومية في جدول DailyObservations.",
        "إتمام وتوثيق 813 قائمة تفتيش يومية في DailySafetyCheckList.",
        "واجهة مخصصة للهواتف الذكية تعمل بالكامل بدون إنترنت (Offline-First).",
        "لوحة متابعة مركزية توضح الملاحظات المفتوحة والمغلقة ونسب الإنجاز.",
        "إلزام المسؤول بإرفاق صورة الإجراء التصحيحي للتحقق قبل إغلاق الملاحظة."
    ], "المرونة: 3,113 ملاحظة و 813 قائمة فحص سُجلت ميدانياً بنجاح", [
        "3,113 ملاحظة سلامة مسجلة في DailyObservations.",
        "813 قائمة فحص يومية مكتملة في DailySafetyCheckList.",
        "عمل كامل في المناطق المعزولة بدون انقطاع مع مزامنة فورية."
    ], is_ar=True, accent_color=ROYAL_BLUE)

    # SLIDE 17: Sustainability & MoC AR
    s17 = prs.slides.add_slide(blank_layout)
    add_header_footer(s17, "الاستدامة والحوكمة", "الاستدامة البيئية وإدارة التغيير والميزانية", "أبعاد استراتيجية لحماية الموارد والتحكم في التغييرات وترشيد الإنفاق", 17, TOTAL_SLIDES, is_ar=True)
    add_visual_content(s17, "17_sustainability_realistic.png", "أركان الاستدامة وإدارة التغيير", [
        "إدارة استهلاك الموارد: المياه، الكهرباء، والغاز الطبيعي بسجلات مخصصة.",
        "حوكمة 13 صنفاً ونوعاً من المخلفات في WasteManagement_RegularWasteTyp.",
        "إدارة طلبات التغيير التشغيلي في ChangeRequests لمنع المخاطر غير المدروسة.",
        "متابعة بنود الميزانية وطلبات الشراء في SafetyBudgetPurchaseOrders.",
        "تقارير استدامة معتمدة تدعم متطلبات التصدير للأسواق العالمية (EU Green Deal)."
    ], "الاستدامة: حوكمة 13 صنف مخلفات وإدارة شاملة للمياه والكهرباء والغاز", [
        "13 صنف مخلفات معتمد في WasteManagement_RegularWasteTyp.",
        "سجلات دورية لاستهلاك الكهرباء والمياه والغاز الطبيعي.",
        "إدارة طلبات التغيير التشغيلي في ChangeRequests."
    ], is_ar=True, accent_color=PURPLE)

    # SLIDE 18: AI Assistant AR
    s18 = prs.slides.add_slide(blank_layout)
    add_header_footer(s18, "الذكاء الاصطناعي", "مساعد الذكاء الاصطناعي التوليدي (Gemini AI)", "تسخير أحدث نماذج الذكاء الاصطناعي لدعم اتخاذ القرار وتحليل الحوادث فورياً", 18, TOTAL_SLIDES, is_ar=True)
    add_visual_content(s18, "18_ai_assistant_realistic.png", "إمكانات مساعد السلامة الذكي", [
        "تسجيل 101 استشارة وتحليل ذكي موثقة في جدول UserAILog.",
        "مستشار سلامة افتراضي يقدم إجابات فورية وفق معايير OSHA و NFPA و ISO.",
        "التحليل الذكي لتقارير الحوادث واستنباط الأسباب الجذرية واقتراح خطط CAPA.",
        "توليد ومراجعة مسودات تقييم المخاطر (JHA) وإجراءات التشغيل القياسية.",
        "صياغة خطط الطوارئ وسيناريوهات الإخلاء والتقارير التنفيذية بأسلوب احترافي."
    ], "الابتكار: 101 استشارة وتحليل ذكي نفذها محرك Gemini AI بنجاح", [
        "101 عملية استشارة وتحليل مسجلة في UserAILog.",
        "دعم كامل للمحادثة التفاعلية باللغة العربية والإنجليزية.",
        "حماية وخصوصية تامة لبيانات وتقارير الشركة الداخلية."
    ], is_ar=True, accent_color=PURPLE)

    # SLIDE 19: Summary AR
    s19 = prs.slides.add_slide(blank_layout)
    add_header_footer(s19, "الخاتمة والتواصل", "ملخص الأثر التنفيذي وخاتمة العرض التقديمي", "شراكة دائمة نحو بيئة عمل آمنة، مستدامة، وخالية من الحوادث", 19, TOTAL_SLIDES, is_ar=True)
    add_visual_content(s19, "03_dashboard_realistic.png", "أبرز مكتسبات تطبيق SafetyHub بالأرقام", [
        "1,654 تصريح عمل رقمي تم إنجازه بنسبة حوكمة 100% وبلا أوراق.",
        "3,113 ملاحظة ميدانية و 813 قائمة فحص عولجت عبر الهواتف الذكية.",
        "تدريب 2,763 عاملاً ومقاولاً لترسيخ ثقافة السلامة الاستباقية.",
        "متابعة وإغلاق 262 إجراءً تصحيحياً في سجل المتابعة الموحد (ATR).",
        "تأمين كامل للمنشأة مدعوم بـ 41,744 سجل تدقيق أمني وحماية سيبرانية."
    ], "الخاتمة: SafetyHub | ICAPP — المنظومة الرائدة للسلامة والصحة والبيئة", [
        "مسؤول النظام: م. ياسر دياب (Yasser Diab).",
        "البريد الإلكتروني: Yasser.diab@icapp.com.eg",
        "دعم فني مباشر وتطوير مستمر بالتعاون مع HSEHub360."
    ], is_ar=True, accent_color=ROYAL_BLUE)

    out_ar = os.path.join(PROJECT_ROOT, "SafetyHub_ICAPP_Visual_Summary_AR.pptx")
    prs.save(out_ar)
    print(f"Arabic presentation saved: {out_ar} ({len(prs.slides)} slides)")

def generate_english_presentation():
    print("Generating English Presentation with REAL Database Numbers...")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    TOTAL_SLIDES = 19

    # SLIDE 1: Cover EN
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY_DARK
    bg1.line.fill.background()

    acc1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.6), Inches(12.133), Inches(6.3))
    acc1.fill.background()
    acc1.line.color.rgb = ROYAL_BLUE
    acc1.line.width = Pt(1.5)

    if os.path.exists(LOGO_PATH):
        try:
            s1.shapes.add_picture(LOGO_PATH, Inches(5.66), Inches(1.1), height=Inches(1.2))
        except Exception:
            pass

    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(2.45), Inches(11.333), Inches(1.1))
    p1 = t_box.text_frame.paragraphs[0]
    p1.text = "SafetyHub | ICAPP"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)
    p1.font.name = "Segoe UI"

    s_box = s1.shapes.add_textbox(Inches(1.0), Inches(3.55), Inches(11.333), Inches(0.6))
    ps1 = s_box.text_frame.paragraphs[0]
    ps1.text = "Executive Visual Presentation & Live Database Walkthrough"
    ps1.alignment = PP_ALIGN.CENTER
    ps1.font.size = Pt(21)
    ps1.font.bold = True
    ps1.font.color.rgb = RGBColor(125, 211, 252)
    ps1.font.name = "Segoe UI"

    s_box2 = s1.shapes.add_textbox(Inches(1.0), Inches(4.25), Inches(11.333), Inches(0.5))
    ps2 = s_box2.text_frame.paragraphs[0]
    ps2.text = "International Company for Agricultural Production & Processing (ICAPP) • Real Operational Data"
    ps2.alignment = PP_ALIGN.CENTER
    ps2.font.size = Pt(15)
    ps2.font.color.rgb = RGBColor(226, 232, 240)
    ps2.font.name = "Segoe UI"

    badge_titles = [
        "📊 1,654 Digital PTWs Issued",
        "📋 3,113 Field Observations Logged",
        "🏥 1,029 Clinic Visits Tracked",
        "🔒 41,744 Security Logs Audited"
    ]
    bw = Inches(2.7)
    b_gap = Inches(0.25)
    start_bx = Inches(1.0)
    for i, b_text in enumerate(badge_titles):
        bx = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, start_bx + i * (bw + b_gap), Inches(5.1), bw, Inches(0.55))
        bx.fill.solid()
        bx.fill.fore_color.rgb = NAVY_MID
        bx.line.color.rgb = ROYAL_BLUE
        bx.line.width = Pt(1)
        p_bx = bx.text_frame.paragraphs[0]
        p_bx.text = b_text
        p_bx.alignment = PP_ALIGN.CENTER
        p_bx.font.size = Pt(11)
        p_bx.font.bold = True
        p_bx.font.color.rgb = RGBColor(255, 255, 255)
        p_bx.font.name = "Segoe UI"

    f_box = s1.shapes.add_textbox(Inches(1.0), Inches(6.0), Inches(11.333), Inches(0.5))
    p_f1 = f_box.text_frame.paragraphs[0]
    p_f1.text = "HSE & Environmental Management Department — ICAPP • Production Data Validated © 2026"
    p_f1.alignment = PP_ALIGN.CENTER
    p_f1.font.size = Pt(11)
    p_f1.font.color.rgb = RGBColor(148, 163, 184)
    p_f1.font.name = "Segoe UI"

    # SLIDE 2: Login EN
    s2 = prs.slides.add_slide(blank_layout)
    add_header_footer(s2, "Authentication & Security", "Secure Login Portal & Two-Factor Authentication (2FA)", "Enterprise-grade authentication with SHA-256 encryption & role-based access", 2, TOTAL_SLIDES, is_ar=False)
    add_visual_content(s2, "02_login_2fa_realistic.png", "Login Portal Key Capabilities", [
        "Secure institutional login with SHA-256 password hashing.",
        "Two-Factor Authentication (2FA/MFA) via Google Authenticator.",
        "Quick PIN-code access option for fast on-site field inspections.",
        "Instant Arabic / English toggle with automatic RTL/LTR adjustment.",
        "Comprehensive cyber audit trail with 41,744 logged security events."
    ], "Security: Robust cyber defense safeguarding corporate assets", [
        "41,744 security audit records in SecurityAuditLog.",
        "6,470 user activity transactions in UserActivityLog.",
        "Automatic session timeout and encrypted token storage."
    ], is_ar=False, accent_color=ROYAL_BLUE)

    # SLIDE 3: Dashboard EN
    s3 = prs.slides.add_slide(blank_layout)
    add_header_footer(s3, "Executive Dashboard", "Real-Time Executive Safety Analytics & KPIs", "24/7 centralized operational visibility into health, safety, and environment metrics", 3, TOTAL_SLIDES, is_ar=False)
    add_visual_content(s3, "03_dashboard_realistic.png", "Dashboard Core Features", [
        "Live metric cards: Safe working man-hours, active permits, and open items.",
        "Tracking 57 official Safety Performance KPIs (TRIR, LTIFR, Severity).",
        "Interactive analytics mapping risks across 132 registered factory locations.",
        "Smart alert center monitoring critical conditions and overdue inspections.",
        "One-click quick actions to initiate permits, log observations, or report hazards."
    ], "Impact: Real-time intelligence empowering proactive hazard elimination", [
        "57 official HSE KPIs monitored in SafetyPerformanceKPIs.",
        "132 distinct operational facility locations in Form_Places.",
        "Live reactive data stream without manual page refresh."
    ], is_ar=False, accent_color=EMERALD)

    # SLIDE 4: PTW EN
    s4 = prs.slides.add_slide(blank_layout)
    add_header_footer(s4, "Critical Operations", "Electronic Permit to Work System (e-PTW)", "Strict digital governance over high-risk hazardous industrial operations", 4, TOTAL_SLIDES, is_ar=False)
    add_visual_content(s4, "04_ptw_realistic.png", "Permit to Work Workflow", [
        "1,654 digital work permits issued and archived in the active PTW database.",
        "1,941 lifecycle transaction records logged in the master PTWRegistry.",
        "Authorized signatory network: 63 internal authorities & 29 contractor authorities.",
        "Coverage: Hot Work, Confined Space, Working at Height, and Electrical LOTO.",
        "Encrypted digital signatures with instant QR Code validation for site patrols."
    ], "Outcome: 1,654 digital permits successfully executed with zero severe incidents", [
        "1,654 electronic permits recorded in PTW table.",
        "1,941 transactions in PTWRegistry master log.",
        "63 internal + 29 contractor certified issuing authorities."
    ], is_ar=False, accent_color=CRIMSON)

    # SLIDE 5: Incidents EN
    s5 = prs.slides.add_slide(blank_layout)
    add_header_footer(s5, "Incident Investigation", "Incident Management & Root Cause Analysis (RCA)", "Systematic incident reporting and deep RCA preventing recurrence", 5, TOTAL_SLIDES, is_ar=False)
    add_visual_content(s5, "05_incidents_realistic.png", "Incident Management Process", [
        "Instant field logging with evidence photo upload and injury classification.",
        "Integrated investigation tools: 5-Whys methodology and Ishikawa Fishbone.",
        "Deep causal factor analysis: Human, Machine, Method, Environment, Material.",
        "Connected to the central Action Tracking Register (262 corrective actions).",
        "Automated Safety Alerts dissemination to prevent similar events across sites."
    ], "Impact: Root causes resolved permanently to eliminate injury recurrence", [
        "Direct integration with central ATR tracking 262 actions.",
        "Direct link between incident records and medical clinic.",
        "Compliant documentation for labor authorities & insurers."
    ], is_ar=False, accent_color=CRIMSON)

    # SLIDE 6: Near-Miss EN
    s6 = prs.slides.add_slide(blank_layout)
    add_header_footer(s6, "Proactive Prevention", "Near-Miss Reporting & Hazard Identification", "Capturing early warning signals to eliminate hazards before incidents occur", 6, TOTAL_SLIDES, is_ar=False)
    add_visual_content(s6, "06_nearmiss_realistic.png", "Near-Miss Reporting System", [
        "Simplified mobile interface encouraging frontline employees to report freely.",
        "Potential risk severity categorization for prioritized corrective actions.",
        "Confidential reporting option building trust and transparent reporting culture.",
        "Departmental trend analytics to focus protective resources effectively.",
        "Recognition programs rewarding proactive employees identifying severe hazards."
    ], "Philosophy: Every resolved near-miss is a serious accident prevented", [
        "Direct assignment of near-miss items into supervisor tasks.",
        "Direct camera capture attachment from mobile devices.",
        "Real-time SLA tracking for on-time hazard resolution."
    ], is_ar=False, accent_color=AMBER)

    # SLIDE 7: Clinic EN
    s7 = prs.slides.add_slide(blank_layout)
    add_header_footer(s7, "Occupational Health", "Medical Clinic & Health Surveillance System", "Secure digital medical records and employee health monitoring programs", 7, TOTAL_SLIDES, is_ar=False)
    add_visual_content(s7, "07_clinic_realistic.png", "Medical Clinic Capabilities", [
        "1,029 documented clinic visits recorded in ClinicContractorVisits.",
        "573 medical dispensing transactions tracked in MedicationDispenseLog.",
        "187 medical staff attendance shifts recorded in ClinicStaffAttendance.",
        "Documentation of sick leaves, occupational treatments, and lost workdays.",
        "Periodic fitness examinations for food handlers and chemical operators."
    ], "Care: 1,029 clinic visits & 573 medication dispenses tracked digitally", [
        "1,029 clinic examinations in ClinicContractorVisits.",
        "573 medication dispenses in MedicationDispenseLog.",
        "187 healthcare shifts logged in ClinicStaffAttendance."
    ], is_ar=False, accent_color=EMERALD)

    # SLIDE 8: Fire Equipment EN
    s8 = prs.slides.add_slide(blank_layout)
    add_header_footer(s8, "Fire Protection", "Fire Fighting & Emergency Equipment Management", "Guaranteed 100% operational readiness for all fire suppression assets", 8, TOTAL_SLIDES, is_ar=False)
    add_visual_content(s8, "08_fire_equipment_realistic.png", "Fire Assets Governance", [
        "169 certified fire protection assets registered in FireEquipmentAssets.",
        "26 formal periodic inspection cycles completed in FireEquipmentInspections.",
        "High-speed field inspections via mobile QR Code scanning on every unit.",
        "Verification of pressure gauges, seal integrity, discharge hoses, and powder.",
        "Automated alerts for periodic hydrostatic pressure testing and cylinder refilling."
    ], "Readiness: 169 fire assets inspected and mapped across the complex", [
        "169 firefighting equipment units in FireEquipmentAssets.",
        "26 complete inspection cycles in FireEquipmentInspections.",
        "100% compliance with Civil Defense emergency standards."
    ], is_ar=False, accent_color=CRIMSON)

    # SLIDE 9: Contractors EN
    s9 = prs.slides.add_slide(blank_layout)
    add_header_footer(s9, "Partners & Site Security", "Contractors Management & Gate Security", "External workforce governance and rigorous perimeter access control", 9, TOTAL_SLIDES, is_ar=False)
    add_visual_content(s9, "09_contractors_realistic.png", "Contractor & Gate Management", [
        "1,736 contractor workers inducted and certified in ContractorTrainings.",
        "194 contractor approval requests & 56 performance evaluation requests.",
        "Blacklist register containing 13 non-compliant vendors in Blacklist_Register.",
        "29 gate visitor & commercial vehicle movements recorded in GateVisitors.",
        "Live headcount tracking showing exactly who is inside the facility 24/7."
    ], "Governance: 1,736 contractor staff trained & 194 approval requests governed", [
        "1,736 contractor workers certified in ContractorTrainings.",
        "194 formal contractor approvals in ContractorApprovalRequests.",
        "13 non-compliant parties blacklisted to protect site safety."
    ], is_ar=False, accent_color=ROYAL_BLUE)

    # SLIDE 10: PPE EN
    s10 = prs.slides.add_slide(blank_layout)
    add_header_footer(s10, "Personal Protection", "PPE Management & Digital Inventory System", "The vital last line of defense with rationalized allocation and stock control", 10, TOTAL_SLIDES, is_ar=False)
    add_visual_content(s10, "10_ppe_realistic.png", "PPE Governance & Distribution", [
        "162 PPE inventory items managed with 105 transactions in PPE_Transactions.",
        "Role-specific PPE requirement matrices covering 12 departments in PPEMatrix.",
        "Enforced lifespan policies preventing redundant or unjustified re-issuance.",
        "Real-time warehouse stock tracking with minimum reorder threshold alerts.",
        "Digital employee sign-off acknowledging receipt and mandatory daily wear."
    ], "Optimization: 162 equipment items & 105 issuance logs tracked", [
        "162 PPE items recorded in the central PPE catalog.",
        "105 issuance transactions in PPE_Transactions.",
        "12 job hazard eligibility matrices in PPEMatrix."
    ], is_ar=False, accent_color=AMBER)

    # SLIDE 11: Chemical Safety EN
    s11 = prs.slides.add_slide(blank_layout)
    add_header_footer(s11, "Hazardous Materials", "Chemical Safety & Safety Data Sheets (MSDS/SDS)", "Comprehensive hazard control, safe storage, and chemical spill readiness", 11, TOTAL_SLIDES, is_ar=False)
    add_visual_content(s11, "11_chemical_safety_realistic.png", "Chemical Safety Protocols", [
        "Central inventory of all chemicals registered in Chemical_Register with CAS numbers.",
        "Bilingual digital repository of Material Safety Data Sheets (MSDS) with quick search.",
        "Chemical incompatibility matrix preventing hazardous co-storage reactions.",
        "Instant emergency medical protocols and mandatory specialized PPE directives.",
        "Spill containment strategy with secondary containment and specialized Spill Kits."
    ], "Protection: Safeguarding personnel and agricultural environment from toxicity", [
        "Formal chemical register managed in Chemical_Register.",
        "GHS standardized visual hazard pictograms and labeling.",
        "Weekly verification of emergency eye-wash and safety shower stations."
    ], is_ar=False, accent_color=ROYAL_BLUE)

    # SLIDE 12: Periodic Inspections EN
    s12 = prs.slides.add_slide(blank_layout)
    add_header_footer(s12, "Periodic Inspections", "Periodic Inspections & Equipment Integrity", "Preventive maintenance ensuring integrity of machinery, cranes, and electrics", 12, TOTAL_SLIDES, is_ar=False)
    add_visual_content(s12, "12_periodic_inspections_realistic.png", "Inspection Scope & Execution", [
        "Scheduled inspections logged in PeriodicEquipmentInspections.",
        "Standardized smart checklists aligned with international engineering codes.",
        "Mandatory photographic evidence for any defective or non-compliant item.",
        "Automated translation of identified defects into maintenance work orders.",
        "Asset historical maintenance record tracking recurrent failures and costs."
    ], "Reliability: Eliminating catastrophic breakdowns and maintaining plant uptime", [
        "Formal machinery inspection records in PeriodicInspections.",
        "Scheduled checkups for forklifts, boilers, and power distribution.",
        "Automated work order dispatch to maintenance engineers."
    ], is_ar=False, accent_color=EMERALD)

    # SLIDE 13: Violations EN
    s13 = prs.slides.add_slide(blank_layout)
    add_header_footer(s13, "Disciplinary Actions", "Safety Violations Management & Enforced Policies", "Fair and firm enforcement of safety rules fostering an accountable culture", 13, TOTAL_SLIDES, is_ar=False)
    add_visual_content(s13, "13_violations_realistic.png", "Violation Governance Workflow", [
        "154 verified safety violation incidents recorded in Violations table.",
        "73 violation approval and review requests processed in ViolationApprovalRequests.",
        "36 standardized violation types cataloged in ViolationTypes.",
        "Transparent progressive disciplinary ladder: Verbal -> Written -> Suspension -> Penalty.",
        "Integration of violation history into annual staff and contractor scorecards."
    ], "Accountability: 154 violations & 73 formal reviews resolved transparently", [
        "154 field violation incidents logged in Violations.",
        "73 administrative reviews in ViolationApprovalRequests.",
        "36 standard violation classifications in ViolationTypes."
    ], is_ar=False, accent_color=AMBER)

    # SLIDE 14: Risk Assessment EN
    s14 = prs.slides.add_slide(blank_layout)
    add_header_footer(s14, "Risk Management", "Risk Assessment, 5x5 Matrix & SOP/JHA Library", "Proactive hazard identification applying the hierarchy of controls", 14, TOTAL_SLIDES, is_ar=False)
    add_visual_content(s14, "14_risk_assessment_realistic.png", "Risk Assessment System Tools", [
        "Interactive 5x5 Risk Matrix calculating probability and consequence severity.",
        "Governance of 40 legal compliance registers cataloged in LegalDocuments.",
        "Job Hazard Analysis (JHA/JSA) breaking jobs into steps with safe controls.",
        "Standard Operating Procedures (SOP) digital repository for all machinery.",
        "Hierarchy of controls: Elimination -> Substitution -> Engineering -> Admin -> PPE."
    ], "Proactivity: Controlling operational risks & maintaining 40 legal registers", [
        "40 verified statutory compliance registers in LegalDocuments.",
        "Direct linkage between high risk activities and work permits.",
        "Clear visual safety instructions easily understood by frontline staff."
    ], is_ar=False, accent_color=ROYAL_BLUE)

    # SLIDE 15: Training EN
    s15 = prs.slides.add_slide(blank_layout)
    add_header_footer(s15, "Competency & Training", "Safety Training Management & Induction Programs", "Empowering the workforce through structured, measurable safety education", 15, TOTAL_SLIDES, is_ar=False)
    add_visual_content(s15, "15_training_realistic.png", "Training & Knowledge Delivery", [
        "439 internal safety training sessions documented in Training table.",
        "1,027 employee attendances recorded + 1,736 contractor worker trainings.",
        "Competency matrices covering 158 roles in EmployeeTrainingMatrix.",
        "Digital certificates generated with authenticable verification QR codes.",
        "Extensive Toolbox Talks (TBT) library providing daily pre-shift topics."
    ], "Competency: 439 training courses & 2,763 trainees certified in database", [
        "439 in-house training courses in Training table.",
        "1,027 employee records in TrainingAttendance + 1,736 contractors.",
        "158 occupational training profiles in EmployeeTrainingMatrix."
    ], is_ar=False, accent_color=EMERALD)

    # SLIDE 16: Forms Hub EN
    s16 = prs.slides.add_slide(blank_layout)
    add_header_footer(s16, "Field Forms Hub", "Daily Safety Observations & Mobile Forms Hub", "Empowering field supervisors with mobile, offline-capable digital forms", 16, TOTAL_SLIDES, is_ar=False)
    add_visual_content(s16, "16_forms_hub_realistic.png", "Mobile Field Forms Hub", [
        "3,113 daily field safety observations captured in DailyObservations.",
        "813 daily safety checklists completed on-site in DailySafetyCheckList.",
        "Touch-friendly mobile interface operational with 100% offline-first capability.",
        "Centralized observation dashboard monitoring open, in-progress, and closed items.",
        "Mandatory photographic evidence of rectification before observation sign-off."
    ], "Agility: 3,113 observations & 813 inspection checklists logged digitally", [
        "3,113 observation reports in DailyObservations.",
        "813 completed checklists in DailySafetyCheckList.",
        "Zero data loss in remote crop fields with automatic cloud synchronization."
    ], is_ar=False, accent_color=ROYAL_BLUE)

    # SLIDE 17: Sustainability & MoC EN
    s17 = prs.slides.add_slide(blank_layout)
    add_header_footer(s17, "Sustainability & MoC", "Environmental Sustainability, MoC & HSE Budget", "Strategic stewardship protecting resources, governing change, and ROI", 17, TOTAL_SLIDES, is_ar=False)
    add_visual_content(s17, "17_sustainability_realistic.png", "Sustainability & Change Control", [
        "Comprehensive tracking of Water, Electricity, and Natural Gas consumption.",
        "13 regulated waste categories tracked in WasteManagement_RegularWasteTyp.",
        "Management of Change (MoC) requests formally governed in ChangeRequests.",
        "HSE budget purchase orders and allocations in SafetyBudgetPurchaseOrders.",
        "Certified sustainability reporting meeting European export standards."
    ], "Sustainability: 13 regulated waste types & complete utility monitoring", [
        "13 regulated waste streams in WasteManagement_RegularWasteTyp.",
        "Continuous metering of electrical, water, and natural gas consumption.",
        "Formal engineering review in ChangeRequests."
    ], is_ar=False, accent_color=PURPLE)

    # SLIDE 18: AI Assistant EN
    s18 = prs.slides.add_slide(blank_layout)
    add_header_footer(s18, "Artificial Intelligence", "Generative AI Assistant Powered by Gemini AI", "Leveraging state-of-the-art AI for real-time safety reasoning and analysis", 18, TOTAL_SLIDES, is_ar=False)
    add_visual_content(s18, "18_ai_assistant_realistic.png", "AI Assistant Core Strengths", [
        "101 AI analytical queries and consultation sessions in UserAILog.",
        "Instant safety advisor answering queries aligned with OSHA, NFPA, and ISO.",
        "Intelligent incident narrative analysis extracting root causes and proposing CAPA.",
        "Automated drafting and review of Job Hazard Analyses (JHA) and safety SOPs.",
        "Rapid technical translation of complex chemical data and safety standards."
    ], "Innovation: 101 AI safety consultations executed with high precision", [
        "101 AI interactions logged in UserAILog.",
        "Full bilingual support for Arabic and English technical queries.",
        "Strict enterprise privacy ensuring internal reports remain confidential."
    ], is_ar=False, accent_color=PURPLE)

    # SLIDE 19: Summary EN
    s19 = prs.slides.add_slide(blank_layout)
    add_header_footer(s19, "Conclusion & Contacts", "Executive Value Summary & Implementation Impact", "A long-term partnership driving a zero-harm, sustainable workplace", 19, TOTAL_SLIDES, is_ar=False)
    add_visual_content(s19, "03_dashboard_realistic.png", "Major SafetyHub Achievements with Real Data", [
        "1,654 work permits executed 100% digitally with zero severe lost time injuries.",
        "3,113 daily observations & 813 inspection checklists logged paperlessly.",
        "2,763 employees and contractor staff certified through 439 training programs.",
        "262 corrective actions tracked and resolved in the Action Tracking Register.",
        "Enterprise cyber resilience guaranteed by 41,744 logged security audit events."
    ], "SafetyHub | ICAPP — Protecting People, Powering Performance", [
        "System Administrator: Eng. Yasser Diab.",
        "Contact Email: Yasser.diab@icapp.com.eg",
        "Direct Microsoft Teams support & ongoing innovation with HSEHub360."
    ], is_ar=False, accent_color=ROYAL_BLUE)

    out_en = os.path.join(PROJECT_ROOT, "SafetyHub_ICAPP_Visual_Summary_EN.pptx")
    prs.save(out_en)
    print(f"English presentation saved: {out_en} ({len(prs.slides)} slides)")

if __name__ == "__main__":
    generate_arabic_presentation()
    generate_english_presentation()
